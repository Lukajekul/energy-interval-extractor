from openpyxl import Workbook, load_workbook
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo
import csv
from dateutil import tz
import xlsxwriter
from openpyxl.workbook import Workbook

header = ["From", "To", "Measurment", "Value in KWH", "", "MeasurmentPoint:"]


L1 = 0
L2 = 0

WARNING_DICTIONARY  = {}

FIRST_RUN = True
FIRST_PRINT_RUN = True

FULL_LIST = []

# workbookWrite = Workbook()
# sheetWrite = workbookWrite.active


workbook = Workbook()
page = workbook.active
#page.title = "Test"

def read_file(path):
    openingList = []
    with open(path) as file:
        for line in file:
            splitLine = line.split(";")
            openingList.append(splitLine)
        openingList = openingList[1:]
        for row in openingList:
            global L1
            global L2
            global WARNING_DICTIONARY
            if row[5] != "L3":
                if row[5] not in WARNING_DICTIONARY:
                    WARNING_DICTIONARY[row[5]] = 1
                else: 
                    WARNING_DICTIONARY[row[5]] += 1          
            
            valueNumberKWH = float(row[3].replace(".", ""))
            valueNumberKWH /= 1000000

            row.pop(0)
            global MeasurementPoint
            global FIRST_RUN
            if FIRST_RUN:
                MeasurementPoint = row[0]
                FIRST_RUN = False
            row.pop(0)
            

            dateTime = datetime.strptime(row[0], "%Y-%m-%d %H:%M:%S.%f")
            from_zone = tz.gettz('UTC')
            to_zone = tz.tzlocal()
            dateTime = dateTime.replace(tzinfo=from_zone)
            dateTime = dateTime.astimezone(to_zone)


            durationObject = timedelta(minutes=15)
            
            
            toTime = dateTime + durationObject

            fromTime = datetime.strftime(dateTime, "%Y-%m-%d %H:%M")
            toTime = datetime.strftime(toTime, "%Y-%m-%d %H:%M")

            global directionValue
            if row[2] == 'r':
                directionValue = "Recieved"
            elif row[2] == "s":
                directionValue = "Sent"

            fillInputList(fromTime, toTime, row[3], valueNumberKWH, openingList)

def fillInputList(fromTime, toTime, measurment, valueNumberKWH, openingList):
    # global FIRST_PRINT_RUN
    global FULL_LIST

    row = [fromTime, toTime, measurment, valueNumberKWH]
    FULL_LIST.append(row)

    #if FIRST_PRINT_RUN:
    #     row = ["From", , "MeasurmentHethod", "Value in KWH", "", "MeasumentPoint:", MeasurementPoint]
    #     FULL_LIST.append(row)
    #     row = [fromTime, toTime, measurment, valueNumberKWH, "", "Direction:", directionValue]
    #     FULL_LIST.append(row)
    #     FIRST_PRINT_RUN = False
    # else:
    #     row = [fromTime, toTime, measurment, valueNumberKWH]
    #     FULL_LIST.append(row)
    if len(FULL_LIST) == len(openingList):
        sortedList = sorted(FULL_LIST, key=lambda x: x[1])
        #page.insert_rows(0)
        page['A1'] = "From"
        page['B1'] = "To"
        page['C1'] = "MeasurmentHethod"
        page['D1'] = "Value in KWH"
        page['F1'] = "MeasumentPoint:"
        page['G1'] = MeasurementPoint
        write(sortedList)

def write(sortedList):
    warningMessage = "Warning! There have been"
    lenBegining = len(warningMessage)
    for i in WARNING_DICTIONARY:
        warningMessage += f" {WARNING_DICTIONARY[i]} instances of {i},"
    for row in sortedList:
        page.append(row)
    if len(warningMessage) > lenBegining:
        page['F4'] = f"{warningMessage[:-1]}"
    page['F2'] = "Direction:"
    page['G2'] = directionValue
    workbook.save("../test.xlsx")

def main ():
    path = ""
    read_file(path)

if __name__ == '__main__':
    main()